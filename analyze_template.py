import openpyxl

# Load template XLSX
template_file = r'templates\hoc_ba_mau.xlsx'
wb = openpyxl.load_workbook(template_file)

print("=== PHÂN TÍCH FILE TEMPLATE ===\n")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ---")
    
    # Print all data to see structure
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        # Only print non-empty rows
        if any(cell is not None for cell in row):
            print(f"Row {row_idx}: {row}")
