# -*- coding: utf-8 -*-
"""
Tool chuyển PDF học bạ tiểu học sang file Excel theo mẫu.
- Có giao diện chọn thư mục input/output/template
- Grid danh sách PDF, tick chọn, chọn tất cả, xóa list, chuyển các file đã chọn
- Kết quả đặt tên: <Lớp>_<Tên học sinh>.xlsx

Tác giả: Võ Thành Khải - 0913046881 / 0913784333
"""
from __future__ import annotations

import os
import re
import sys
import shutil
import csv
import json
import threading
import traceback
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import fitz  # PyMuPDF
from openpyxl import load_workbook

try:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
except ImportError:
    class _MissingTk:
        Tk = object

    tk = _MissingTk()
    ttk = filedialog = messagebox = None

try:
    from updater import AutoUpdater, VersionManager
    _HAS_UPDATER = True
except Exception:
    _HAS_UPDATER = False


# =========================
# 1) HÀM TIỆN ÍCH
# =========================

def app_base_dir() -> Path:
    """Lấy thư mục chạy app, hỗ trợ cả khi đóng gói PyInstaller."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def norm_text(s: object) -> str:
    if s is None:
        return ""
    s = str(s).replace("\xa0", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def norm_key(s: str) -> str:
    """Chuẩn hóa để so khớp nhãn, vẫn giữ tiếng Việt nhưng bỏ khoảng trắng thừa."""
    return norm_text(s).lower().replace("đ", "d")


def safe_filename(s: str, max_len: int = 120) -> str:
    s = norm_text(s)
    s = re.sub(r"[\\/:*?\"<>|]+", "_", s)
    s = re.sub(r"\s+", "_", s)
    s = s.strip("._ ")
    return s[:max_len] if len(s) > max_len else s


def after_colon(text: str) -> str:
    text = norm_text(text)
    return text.split(":", 1)[1].strip() if ":" in text else text


def regex_get(text: str, pattern: str, default: str = "") -> str:
    m = re.search(pattern, text, flags=re.I | re.M)
    return norm_text(m.group(1)) if m else default


def clean_sentence(s: str) -> str:
    s = norm_text(s)
    return s.strip()


def collapse_comment_lines(lines: List[str]) -> List[str]:
    comments: List[str] = []
    current: List[str] = []
    for line in lines:
        line = clean_sentence(line)
        if not line:
            continue
        current.append(line)
        if re.search(r"[.!?…]$", line):
            comments.append(clean_sentence(" ".join(current)))
            current = []
    if current:
        comments.append(clean_sentence(" ".join(current)))
    return comments


# =========================
# 2) MÔ HÌNH DỮ LIỆU
# =========================

@dataclass
class SubjectResult:
    level: str = ""
    score: str = ""
    comment: str = ""


@dataclass
class ThcsSubjectResult:
    hk1: str = ""
    hk2: str = ""
    year: str = ""
    retake: str = ""
    teacher: str = ""


@dataclass
class StudyRecord:
    school_year: str = ""
    class_name: str = ""
    school: str = ""


@dataclass
class SummaryResult:
    conduct: str = ""
    learning: str = ""
    absent: str = ""
    conduct_retake: str = ""
    learning_retake: str = ""


@dataclass
class HocBaData:
    education_level: str = ""
    student_name: str = ""
    gender: str = ""
    birth_date: str = ""
    ethnicity: str = ""
    nationality: str = ""
    birth_place: str = ""
    hometown: str = ""
    address: str = ""
    father: str = ""
    mother: str = ""
    guardian: str = ""
    school: str = ""
    ward: str = ""
    district: str = ""
    province: str = ""
    class_name: str = ""
    school_year: str = ""
    height: str = ""
    weight: str = ""
    absent_allowed: str = ""
    absent_not_allowed: str = ""
    principal: str = ""
    principal_title: str = ""
    principal_subtitle: str = ""
    teacher: str = ""
    sign_date: str = ""                  # ngày ký (cuối trang kết quả/nhận xét) → [130]
    issue_date: str = ""                 # ngày phát hành học bạ (trang bìa/thông tin) → [19]
    sign_page_principal_title: str = ""  # chức danh HT trang tổng kết → [132]
    entry_date: str = ""
    result: str = ""
    reward: str = ""
    completion: str = ""
    school_register_no: str = ""
    policy_object: str = ""
    father_job: str = ""
    mother_job: str = ""
    guardian_job: str = ""
    promotion: str = ""
    promotion_after_retake: str = ""
    no_promotion: str = ""
    certificate: str = ""
    contest_result: str = ""
    homeroom_comment: str = ""
    subjects: Dict[str, SubjectResult] = field(default_factory=dict)
    qualities: Dict[str, SubjectResult] = field(default_factory=dict)
    common_competencies: Dict[str, SubjectResult] = field(default_factory=dict)
    specific_competencies: Dict[str, SubjectResult] = field(default_factory=dict)
    thcs_subjects: Dict[str, ThcsSubjectResult] = field(default_factory=dict)
    study_records: List[StudyRecord] = field(default_factory=list)
    summary_results: Dict[str, SummaryResult] = field(default_factory=dict)


# =========================
# 3) ĐỌC PDF
# =========================

SUBJECT_ITEMS: List[Tuple[str, List[str]]] = [
    ("Tiếng Việt", ["Tiếng Việt"]),
    ("Toán", ["Toán"]),
    ("Ngoại ngữ Tiếng Anh", ["Ngoại ngữ 1", "Tiếng Anh"]),
    ("Ngoại ngữ Tiếng Anh", ["Ngoại ngữ 1"]),
    ("Lịch sử và Địa lí", ["Lịch sử và Địa lí"]),
    ("Khoa học", ["Khoa học"]),
    ("TH-CN (Công nghệ)", ["TH-CN", "(Công nghệ)"]),
    ("TH-CN (Tin học)", ["TH-CN", "(Tin học)"]),
    ("TH-CN (Tin học)", ["Tin học và Công", "nghệ (Tin học)"]),
    ("TH-CN (Công nghệ)", ["Tin học và Công", "nghệ (Công", "nghệ)"]),
    ("Đạo đức", ["Đạo đức"]),
    ("Tự nhiên và Xã hội", ["Tự nhiên và Xã hội"]),
    ("TN-XH", ["TN-XH"]),
    ("Nghệ thuật (Âm nhạc)", ["Nghệ thuật", "(Âm nhạc)"]),
    ("Nghệ thuật (Mĩ thuật)", ["Nghệ thuật", "(Mĩ thuật)"]),
    ("Giáo dục thể chất", ["Giáo dục thể chất"]),
    ("Hoạt động trải nghiệm", ["Hoạt động trải", "nghiệm"]),
    ("Tiếng dân tộc", ["Tiếng dân tộc"]),
]

QUALITY_ITEMS = ["Yêu nước", "Nhân ái", "Chăm chỉ", "Trung thực", "Trách nhiệm"]
COMMON_COMP_ITEMS = ["Tự chủ và tự học", "Giao tiếp và hợp tác", "Giải quyết vấn đề và sáng tạo"]
SPECIFIC_COMP_ITEMS = ["Ngôn ngữ", "Tính toán", "Khoa học", "Công nghệ", "Tin học", "Thẩm mĩ", "Thể chất"]


def read_pdf_text(pdf_path: Path) -> Tuple[str, List[str]]:
    doc = fitz.open(str(pdf_path))
    pages = [page.get_text("text") for page in doc]
    return "\n".join(pages), pages


def lines_from(text: str) -> List[str]:
    return [norm_text(x) for x in text.splitlines() if norm_text(x)]


def find_sequence(lines: List[str], tokens: List[str], start: int = 0) -> Optional[Tuple[int, int]]:
    """Tìm chuỗi token liên tiếp theo dòng; trả về (start_index, end_index_exclusive)."""
    ntokens = [norm_key(t) for t in tokens]
    nlines = [norm_key(x) for x in lines]
    for i in range(start, len(lines)):
        ok = True
        for j, tok in enumerate(ntokens):
            if i + j >= len(nlines) or nlines[i + j] != tok:
                ok = False
                break
        if ok:
            return i, i + len(tokens)
    return None


def parse_items_by_sequence(lines: List[str], items: List[Tuple[str, List[str]]]) -> Dict[str, SubjectResult]:
    found = []
    for key, toks in items:
        res = find_sequence(lines, toks)
        if res:
            found.append((res[0], res[1], key, toks))
    found.sort(key=lambda x: (x[0], -(x[1] - x[0])))

    # Nếu có nhiều cách nhận diện cùng một nhãn tại cùng vị trí, ưu tiên mẫu dài hơn.
    # Ví dụ lớp 1 có "Ngoại ngữ 1" + "Tiếng Anh", lớp 4 chỉ có "Ngoại ngữ 1".
    filtered = []
    occupied_until = -1
    seen_key_start = set()
    for item in found:
        start, end, key, _ = item
        if start < occupied_until:
            continue
        key_start = (key, start)
        if key_start in seen_key_start:
            continue
        filtered.append(item)
        seen_key_start.add(key_start)
        occupied_until = end
    found = filtered

    out: Dict[str, SubjectResult] = {}
    for idx, (start, end, key, toks) in enumerate(found):
        next_start = found[idx + 1][0] if idx + 1 < len(found) else len(lines)
        block = [x for x in lines[end:next_start] if x]
        if not block:
            continue
        level = ""
        score = ""
        comments: List[str] = []
        k = 0
        if k < len(block) and re.fullmatch(r"[THCĐDC]", block[k].strip(), flags=re.I):
            level = block[k].strip().upper()
            k += 1
        if k < len(block) and re.fullmatch(r"\d+(?:[,.]\d+)?", block[k].strip()):
            score = block[k].strip().replace(",", ".")
            k += 1
        comments = block[k:]
        out[key] = SubjectResult(level=level, score=score, comment=clean_sentence(" ".join(comments)))
    return out


def parse_ordered_section(lines: List[str], items: List[str]) -> Dict[str, SubjectResult]:
    """
    Dùng cho phẩm chất/năng lực: PDF thường đặt nhiều nhận xét trong một cột lớn.
    Hàm này lấy mức đạt được sau từng nhãn và phân phối các dòng nhận xét theo thứ tự.
    """
    positions = []
    for item in items:
        toks = item.split(" và ") if item in ["Tự chủ và tự học", "Giao tiếp và hợp tác"] else [item]
        # xử lý nhãn có nhiều dòng phổ biến
        if item == "Tự chủ và tự học": toks = ["Tự chủ", "và tự học"]
        if item == "Giao tiếp và hợp tác": toks = ["Giao tiếp", "và hợp tác"]
        if item == "Giải quyết vấn đề và sáng tạo": toks = ["Giải quyết", "vấn đề", "và sáng tạo"]
        r = find_sequence(lines, toks)
        if r:
            positions.append((r[0], r[1], item))
    positions.sort(key=lambda x: x[0])

    out: Dict[str, SubjectResult] = {item: SubjectResult() for item in items}
    if not positions:
        return out

    # mức đạt sau từng nhãn
    for idx, (start, end, item) in enumerate(positions):
        next_start = positions[idx + 1][0] if idx + 1 < len(positions) else len(lines)
        block = lines[end:next_start]
        for x in block[:4]:
            if re.fullmatch(r"[TĐDC]", x.strip(), flags=re.I):
                out[item].level = x.strip().upper().replace("D", "Đ")
                break

    # lấy nhận xét: các dòng không phải nhãn/mức/tiêu đề, ưu tiên đoạn sau nhãn đầu tiên và trước nhãn tiếp theo nếu có
    used_label_line_idx = set()
    for start, end, _ in positions:
        used_label_line_idx.update(range(start, end))
    candidate_lines: List[str] = []
    for i, line in enumerate(lines):
        if i in used_label_line_idx:
            continue
        if re.fullmatch(r"[TĐDC]", line.strip(), flags=re.I):
            continue
        if any(norm_key(line) == norm_key(h) for h in ["Phẩm chất", "Năng lực", "Mức đạt", "được", "Nhận xét"]):
            continue
        if line.startswith("2.") or line.startswith("3."):
            continue
        # Giữ cả dòng nối câu rất ngắn như "tập.", "cao.", "bạn bè."
        if len(line) >= 2:
            candidate_lines.append(clean_sentence(line))

    candidates = collapse_comment_lines(candidate_lines)

    present_items = [item for _, _, item in positions]
    # Nếu số dòng nhận xét >= số mục thực sự xuất hiện trong PDF, phân bổ từng dòng theo thứ tự xuất hiện.
    # Trường hợp lớp 1 thường chỉ có 5 năng lực đặc thù, không có Công nghệ/Tin học.
    if len(candidates) >= len(present_items):
        for item, cmt in zip(present_items, candidates[:len(present_items)]):
            out[item].comment = cmt
    elif candidates and present_items:
        # nếu chỉ có một đoạn nhận xét, đưa vào dòng đầu tiên có xuất hiện trong PDF
        out[present_items[0]].comment = " ".join(candidates)
    return out


def split_page_section(pages: List[str], page_index: int, start_marker: str, end_marker: Optional[str] = None) -> List[str]:
    if page_index >= len(pages):
        return []
    text = pages[page_index]
    start = text.find(start_marker)
    if start >= 0:
        text = text[start:]
    if end_marker:
        end = text.find(end_marker)
        if end >= 0:
            text = text[:end]
    return lines_from(text)


def find_page_index(pages: List[str], *markers: str, default: Optional[int] = None) -> Optional[int]:
    norm_markers = [norm_key(m) for m in markers if m]
    for idx, page in enumerate(pages):
        page_key = norm_key(page)
        if all(marker in page_key for marker in norm_markers):
            return idx
    return default


def parse_tieu_hoc_pdf(pdf_path: Path) -> HocBaData:
    text, pages = read_pdf_text(pdf_path)
    data = HocBaData()

    data.student_name = regex_get(text, r"Họ và tên học sinh:[ \t]*([^\n]+)")
    data.gender = regex_get(text, r"Giới tính:[ \t]*([^\n]+)")
    data.birth_date = regex_get(text, r"Ngày, tháng, năm sinh:[ \t]*([^\n]+)")
    data.ethnicity = regex_get(text, r"Dân tộc:[ \t]*([^\n]+)")
    data.nationality = regex_get(text, r"Quốc tịch:[ \t]*([^\n]+)")
    data.birth_place = regex_get(text, r"Nơi sinh:[ \t]*([^\n]*)")
    data.hometown = regex_get(text, r"Quê quán:[ \t]*([^\n]*)")
    data.address = regex_get(text, r"Nơi ở hiện nay:[ \t]*([^\n]+)")
    data.father = regex_get(text, r"Họ và tên cha:[ \t]*([^\n]+)")
    data.mother = regex_get(text, r"Họ và tên mẹ:[ \t]*([^\n]+)")
    data.guardian = regex_get(text, r"Người giám hộ\s*\(?nếu có\)?:[ \t]*([^\n]*)")
    data.school = regex_get(text, r"Trường[ \t]*:?[ \t]*([^\n]+)")
    data.ward = regex_get(text, r"Xã \(Phường[^\n]*\):[ \t]*([^\n]+)")
    data.province = regex_get(text, r"Tỉnh \(Thành phố\):[ \t]*([^\n]+)")
    data.class_name = regex_get(text, r"Lớp:[ \t]*([^\n]+)")
    data.school_year = regex_get(text, r"Năm học\s*([0-9]{4}\s*-\s*[0-9]{4})")
    if not data.school_year:
        data.school_year = regex_get(text, r"([0-9]{4}\s*-\s*[0-9]{4})")
    data.height = regex_get(text, r"Chiều cao:[ \t]*([^\n]+)")
    data.weight = regex_get(text, r"Cân nặng:[ \t]*([^\n]+)")
    data.absent_allowed = regex_get(text, r"Số ngày nghỉ có phép:[ \t]*([^\n]+)")
    data.absent_not_allowed = regex_get(text, r"Số ngày nghỉ không phép:[ \t]*([^\n]+)")
    # Ngày phát hành học bạ: lần xuất hiện đầu tiên của "ngày X tháng Y năm Z" trong toàn bộ text
    data.issue_date = regex_get(text, r"([^\n]*ngày\s+\d{1,2}\s+tháng\s+\d{1,2}\s+năm\s+\d{4})")

    subject_page = find_page_index(pages, "1. Các môn học", "Lớp:", default=3 if len(pages) > 3 else None)
    competency_page = find_page_index(pages, "2. Những phẩm chất", "3.1", default=4 if len(pages) > 4 else None)
    result_page = find_page_index(pages, "4. Đánh giá kết quả giáo dục", "Xác nhận", default=5 if len(pages) > 5 else None)

    # Hiệu trưởng/GVCN + ngày ký: lấy từ trang kết quả
    if result_page is not None:
        p6lines = lines_from(pages[result_page])
        if p6lines:
            # thường 2 dòng tên cuối là hiệu trưởng và giáo viên chủ nhiệm
            name_lines = [x for x in p6lines if re.search(r"^[A-ZÀ-ỸĐ][A-Za-zÀ-ỹđĐ\s]+$", x) and len(x.split()) >= 2]
            if len(name_lines) >= 2:
                data.principal = name_lines[-2]
                data.teacher = name_lines[-1]
            # Ngày ký: lấy dòng cuối cùng có dạng "ngày X tháng Y năm Z" trên trang kết quả
            sign_lines = [x for x in p6lines if re.search(r"ngày\s+\d{1,2}\s+tháng\s+\d{1,2}\s+năm\s+\d{4}", x, re.I)]
            if sign_lines:
                data.sign_date = sign_lines[-1]

    # Mục môn học
    if subject_page is not None:
        p4lines = lines_from(pages[subject_page])
        data.subjects = parse_items_by_sequence(p4lines, SUBJECT_ITEMS)

    # Mục phẩm chất/năng lực
    if competency_page is not None:
        p5 = pages[competency_page]
        q_lines = lines_from(p5[p5.find("2. Những phẩm chất") : p5.find("3. Những năng lực") if "3. Những năng lực" in p5 else len(p5)])
        cc_lines = lines_from(p5[p5.find("3.1") : p5.find("3.2") if "3.2" in p5 else len(p5)])
        sc_lines = lines_from(p5[p5.find("3.2") :])
        data.qualities = parse_ordered_section(q_lines, QUALITY_ITEMS)
        data.common_competencies = parse_ordered_section(cc_lines, COMMON_COMP_ITEMS)
        data.specific_competencies = parse_ordered_section(sc_lines, SPECIFIC_COMP_ITEMS)

    # Trang 6: kết quả, khen thưởng, hoàn thành
    data.result = regex_get(text, r"4\.\s*Đánh giá kết quả giáo dục:\s*([^\n]+)").rstrip(".")
    reward_match = re.search(r"5\.\s*Khen thưởng:\s*\n?(.+?)\n6\.\s*Hoàn thành", text, flags=re.S | re.I)
    if reward_match:
        data.reward = clean_sentence(re.sub(r"\n+", " ", reward_match.group(1)).lstrip("- "))
    completion_match = re.search(r"6\.\s*Hoàn thành chương trình lớp học/chương trình tiểu học:\s*\n?(.+?)(?:\n[^\n]*ngày|$)", text, flags=re.S | re.I)
    if completion_match:
        data.completion = clean_sentence(re.sub(r"\n+", " ", completion_match.group(1)).lstrip("- "))

    # Trang quá trình học tập
    data.entry_date = regex_get(text, r"(\d{2}/\d{2}/\d{4})")

    return data


THCS_SUBJECT_ITEMS: List[Tuple[str, List[str]]] = [
    ("Toán học", ["Toán học"]),
    ("Lịch sử và Địa lí", ["Lịch sử và Địa lí"]),
    ("Khoa học tự nhiên", ["Khoa học tự nhiên"]),
    ("Tin học", ["Tin học"]),
    ("Ngữ văn", ["Ngữ văn"]),
    ("Ngoại ngữ 1 (Tiếng Anh)", ["Ngoại ngữ 1", "(Tiếng Anh)"]),
    ("GDCD", ["GDCD"]),
    ("Công nghệ", ["Công nghệ"]),
    ("Giáo dục thể chất", ["Giáo dục thể chất"]),
    ("Nghệ thuật", ["Nghệ thuật"]),
    ("Ngoại ngữ 2", ["Ngoại ngữ 2", "()"]),
    ("Nội dung giáo dục của địa phương", ["Nội dung giáo dục của", "địa phương"]),
    ("Hoạt động trải nghiệm, hướng nghiệp", ["Hoạt động trải nghiệm,", "hướng nghiệp"]),
]


def clean_pdf_value(value: str) -> str:
    value = norm_text(value)
    value = value.strip(" -")
    if not value or re.fullmatch(r"[.…\.\s/:;-]+", value):
        return ""
    return value


def line_value(lines: List[str], label: str) -> str:
    label_key = norm_key(label)
    for idx, line in enumerate(lines):
        line_key = norm_key(line)
        if line_key.startswith(label_key):
            value = after_colon(line)
            if value == line and idx + 1 < len(lines):
                value = lines[idx + 1]
            if norm_key(label).startswith(norm_key("Đối tượng")):
                value = re.sub(r"^\([^)]*\)\s*", "", value).strip()
            return clean_pdf_value(value)
    return ""


def first_line_after(lines: List[str], marker: str) -> str:
    marker_key = norm_key(marker)
    for idx, line in enumerate(lines):
        if norm_key(line) == marker_key and idx + 1 < len(lines):
            return clean_pdf_value(lines[idx + 1])
    return ""


def is_person_name(line: str) -> bool:
    line = norm_text(line)
    if not line or line.startswith("("):
        return False
    ignored = {
        "HIỆU TRƯỞNG",
        "KT.HIỆU TRƯỞNG",
        "PHÓ HIỆU TRƯỞNG",
        "Giáo viên chủ nhiệm",
        "Hiệu trưởng",
    }
    if line in ignored:
        return False
    return bool(re.search(r"[A-Za-zÀ-ỹĐđ]", line)) and len(line.split()) >= 2


def parse_signature_block(lines: List[str], start_idx: int, end_idx: int) -> Tuple[str, str, str]:
    block = [x for x in lines[start_idx:end_idx] if x and not x.startswith("(")]
    names = [x for x in block if is_person_name(x) and not x.isupper()]
    if not names:
        names = [x for x in block if is_person_name(x)]
    name = names[-1] if names else ""
    titles = [x for x in block if x != name and ("HIỆU TRƯỞNG" in x.upper() or "PHÓ HIỆU TRƯỞNG" in x.upper())]
    title = titles[0] if titles else ""
    subtitle = titles[1] if len(titles) > 1 else ""
    return title, subtitle, name


def parse_study_records(lines: List[str]) -> List[StudyRecord]:
    out: List[StudyRecord] = []
    start = next((i for i, x in enumerate(lines) if norm_key(x) == norm_key("QUÁ TRÌNH HỌC TẬP")), -1)
    if start < 0:
        return out
    tail = lines[start + 1 :]
    i = 0
    while i < len(tail):
        if re.fullmatch(r"\d{4}\s*-\s*\d{4}", tail[i]):
            school_year = tail[i]
            class_name = tail[i + 1] if i + 1 < len(tail) else ""
            school = tail[i + 2] if i + 2 < len(tail) else ""
            out.append(StudyRecord(school_year=school_year, class_name=class_name, school=school))
            i += 3
            continue
        i += 1
    return out


def is_thcs_grade_value(line: str) -> bool:
    value = norm_text(line).replace(",", ".")
    return bool(re.fullmatch(r"\d+(?:\.\d+)?|Đ|CĐ|T|K|Đạt|Chưa đạt", value, flags=re.I))


def parse_thcs_subjects(lines: List[str]) -> Dict[str, ThcsSubjectResult]:
    found = []
    for key, toks in THCS_SUBJECT_ITEMS:
        res = find_sequence(lines, toks)
        if res:
            found.append((res[0], res[1], key))
    found.sort(key=lambda x: x[0])

    out: Dict[str, ThcsSubjectResult] = {}
    footer_markers = (
        "Trong trang này",
        "Xác nhận",
        "(Ký",
        "Ký và ghi rõ",
    )
    for idx, (start, end, key) in enumerate(found):
        next_start = found[idx + 1][0] if idx + 1 < len(found) else len(lines)
        block = lines[end:next_start]
        values: List[str] = []
        teacher_parts: List[str] = []
        for raw in block:
            line = clean_pdf_value(raw)
            if not line:
                continue
            if any(line.startswith(marker) for marker in footer_markers):
                break
            if len(values) < 4 and is_thcs_grade_value(line):
                values.append(line.replace(",", "."))
            elif values:
                teacher_parts.append(line)
        out[key] = ThcsSubjectResult(
            hk1=values[0] if len(values) > 0 else "",
            hk2=values[1] if len(values) > 1 else "",
            year=values[2] if len(values) > 2 else "",
            retake=values[3] if len(values) > 3 else "",
            teacher=clean_sentence(" ".join(teacher_parts)),
        )
    return out


def parse_summary_rows(lines: List[str]) -> Dict[str, SummaryResult]:
    out: Dict[str, SummaryResult] = {}
    labels = {
        "Học kì I": "hk1",
        "Học kì II": "hk2",
        "Cả năm": "year",
    }
    stop_labels = {norm_key(x) for x in labels}
    for label, key in labels.items():
        idx = next((i for i, x in enumerate(lines) if norm_key(x) == norm_key(label)), -1)
        if idx < 0:
            continue
        vals: List[str] = []
        j = idx + 1
        while j < len(lines):
            current = norm_key(lines[j])
            if current in stop_labels or current.startswith(norm_key("Nếu là lớp cuối cấp")):
                break
            value = clean_pdf_value(lines[j])
            if value:
                vals.append(value)
            if len(vals) >= 5:
                break
            j += 1
        out[key] = SummaryResult(
            conduct=vals[0] if len(vals) > 0 else "",
            learning=vals[1] if len(vals) > 1 else "",
            absent=vals[2] if len(vals) > 2 else "",
            conduct_retake=vals[3] if len(vals) > 3 else "",
            learning_retake=vals[4] if len(vals) > 4 else "",
        )
    return out


def collect_after_marker(lines: List[str], marker: str, stop_prefixes: Tuple[str, ...]) -> str:
    marker_key = norm_key(marker)
    for idx, line in enumerate(lines):
        if norm_key(line).startswith(marker_key):
            first = after_colon(line)
            parts: List[str] = []
            if first != line:
                first = clean_pdf_value(first)
                if first:
                    parts.append(first)
            j = idx + 1
            if not parts and marker_key.startswith(norm_key("- Được lên lớp sau")):
                while j < len(lines):
                    candidate_key = norm_key(lines[j])
                    if any(candidate_key.startswith(norm_key(prefix)) for prefix in stop_prefixes):
                        break
                    if lines[j].rstrip().endswith(":"):
                        tail = clean_pdf_value(after_colon(lines[j]))
                        if tail:
                            parts.append(tail)
                        j += 1
                        break
                    j += 1
            while j < len(lines):
                candidate = lines[j]
                candidate_key = norm_key(candidate)
                if any(candidate_key.startswith(norm_key(prefix)) for prefix in stop_prefixes):
                    break
                if marker_key.startswith(norm_key("- Được lên lớp sau")) and (
                    "học hoặc rèn luyện" in candidate or candidate.rstrip().endswith("HK:")
                ):
                    j += 1
                    continue
                value = clean_pdf_value(candidate)
                if value:
                    parts.append(value)
                j += 1
            return clean_sentence(" ".join(parts))
    return ""


def parse_homeroom_comment(lines: List[str]) -> str:
    start = next((i for i, x in enumerate(lines) if norm_key(x).startswith(norm_key("NHẬN XÉT CỦA GIÁO VIÊN CHỦ NHIỆM"))), -1)
    end = next((i for i, x in enumerate(lines) if norm_key(x).startswith(norm_key("Giáo viên chủ nhiệm"))), -1)
    if start < 0 or end < 0 or end <= start:
        return ""
    parts: List[str] = []
    for line in lines[start + 1 : end]:
        key = norm_key(line)
        if (
            key.startswith("(")
            or line.startswith("(")
            or "Ghi nhận xét" in line
            or "về kết quả rèn luyện" in line
            or "và học tập" in line
            or "những vấn đề cần quan tâm" in line
        ):
            continue
        value = clean_pdf_value(line)
        if value:
            parts.append(value)
    return clean_sentence(" ".join(parts))


def parse_thcs_pdf(pdf_path: Path) -> HocBaData:
    text, pages = read_pdf_text(pdf_path)
    data = HocBaData(education_level="thcs")

    cover_lines = lines_from(pages[0]) if pages else []
    data.school = next((line for line in cover_lines if norm_key(line).startswith(norm_key("TRƯỜNG"))), "")
    data.ward = line_value(cover_lines, "Xã/Phường/Đặc khu")
    data.province = line_value(cover_lines, "Tỉnh/Thành phố")
    data.student_name = first_line_after(cover_lines, "Họ và tên học sinh")
    data.school_register_no = line_value(cover_lines, "Số sổ đăng bộ PCGD")

    info_page = find_page_index(pages, "Họ và tên:", "QUÁ TRÌNH HỌC TẬP", default=3 if len(pages) > 3 else None)
    if info_page is not None:
        info_lines = lines_from(pages[info_page])
        data.student_name = line_value(info_lines, "Họ và tên") or data.student_name
        data.gender = line_value(info_lines, "Giới tính")
        data.birth_date = line_value(info_lines, "Ngày sinh")
        data.birth_place = line_value(info_lines, "Nơi sinh")
        data.ethnicity = line_value(info_lines, "Dân tộc")
        data.policy_object = line_value(info_lines, "Đối tượng")
        data.address = line_value(info_lines, "Chỗ ở hiện tại")
        data.father = line_value(info_lines, "Họ và tên cha")
        data.mother = line_value(info_lines, "Họ và tên mẹ")
        data.guardian = line_value(info_lines, "Họ và tên người giám hộ")

        job_values = [after_colon(x) for x in info_lines if norm_key(x).startswith(norm_key("Nghề nghiệp"))]
        if len(job_values) > 0:
            data.father_job = clean_pdf_value(job_values[0])
        if len(job_values) > 1:
            data.mother_job = clean_pdf_value(job_values[1])
        if len(job_values) > 2:
            data.guardian_job = clean_pdf_value(job_values[2])

        # Ngày phát hành học bạ: lấy từ trang thông tin học sinh
        data.issue_date = next((x for x in info_lines if re.search(r"ngày\s+\d{1,2}\s+tháng\s+\d{1,2}\s+năm\s+\d{4}", x, re.I)), "")
        date_idx = info_lines.index(data.issue_date) if data.issue_date in info_lines else -1
        process_idx = next((i for i, x in enumerate(info_lines) if norm_key(x) == norm_key("QUÁ TRÌNH HỌC TẬP")), len(info_lines))
        if date_idx >= 0:
            data.principal_title, data.principal_subtitle, data.principal = parse_signature_block(info_lines, date_idx + 1, process_idx)
        data.study_records = parse_study_records(info_lines)
        if data.study_records:
            data.school_year = data.study_records[-1].school_year
            data.class_name = data.study_records[-1].class_name

    subject_page = find_page_index(pages, "Môn học/Hoạt động", "Học kì I", default=5 if len(pages) > 5 else None)
    if subject_page is not None:
        subject_lines = lines_from(pages[subject_page])
        data.student_name = line_value(subject_lines, "Họ và tên") or data.student_name
        data.class_name = line_value(subject_lines, "Lớp") or data.class_name
        data.school_year = line_value(subject_lines, "Năm học") or data.school_year
        data.thcs_subjects = parse_thcs_subjects(subject_lines)
        names = [x for x in subject_lines if is_person_name(x)]
        if len(names) >= 2:
            data.teacher = names[-2]
            data.principal = data.principal or names[-1]

    summary_page = find_page_index(pages, "Họ tên học sinh", "Mức đánh giá", default=6 if len(pages) > 6 else None)
    if summary_page is not None:
        summary_lines = lines_from(pages[summary_page])
        data.student_name = line_value(summary_lines, "Họ tên học sinh") or data.student_name
        data.class_name = line_value(summary_lines, "Lớp") or data.class_name
        data.school_year = line_value(summary_lines, "Năm học") or data.school_year
        data.summary_results = parse_summary_rows(summary_lines)
        data.promotion = collect_after_marker(summary_lines, "- Được lên lớp", ("- Được lên lớp sau", "- Không được lên lớp", "Kết quả"))
        data.promotion_after_retake = collect_after_marker(summary_lines, "- Được lên lớp sau", ("- Không được lên lớp", "Kết quả"))
        data.no_promotion = collect_after_marker(summary_lines, "- Không được lên lớp", ("Kết quả", "Học kì I"))
        data.completion = collect_after_marker(summary_lines, "Nếu là lớp cuối cấp", ("- Chứng chỉ", "- Kết quả", "- Khen thưởng", "KẾT QUẢ"))
        data.certificate = collect_after_marker(summary_lines, "- Chứng chỉ", ("- Kết quả", "- Khen thưởng", "KẾT QUẢ"))
        data.contest_result = collect_after_marker(summary_lines, "- Kết quả tham gia", ("- Khen thưởng", "KẾT QUẢ"))
        data.reward = collect_after_marker(summary_lines, "- Khen thưởng", ("KẾT QUẢ", "NHẬN XÉT"))
        data.homeroom_comment = parse_homeroom_comment(summary_lines)

        teacher_idx = next((i for i, x in enumerate(summary_lines) if norm_key(x).startswith(norm_key("Giáo viên chủ nhiệm"))), -1)
        if teacher_idx >= 0:
            teacher_candidates = [x for x in summary_lines[teacher_idx + 1 :] if is_person_name(x)]
            if teacher_candidates:
                data.teacher = teacher_candidates[0]
        date_lines = [x for x in summary_lines if re.search(r"ngày\s+\d{1,2}\s+tháng\s+\d{1,2}\s+năm\s+\d{4}", x, re.I)]
        if date_lines:
            data.sign_date = date_lines[-1]
        if data.sign_date in summary_lines:
            date_idx = summary_lines.index(data.sign_date)
            sp_title, _, principal = parse_signature_block(summary_lines, date_idx + 1, len(summary_lines))
            data.principal = principal or data.principal
            data.sign_page_principal_title = sp_title or "Hiệu trưởng"

    return data


def parse_hoc_ba_pdf(pdf_path: Path) -> HocBaData:
    text, _ = read_pdf_text(pdf_path)
    if "TRUNG HỌC CƠ SỞ" in text or "Môn học/Hoạt động" in text:
        return parse_thcs_pdf(pdf_path)
    return parse_tieu_hoc_pdf(pdf_path)


# =========================
# 4) GHI VÀO EXCEL MẪU
# =========================

SUBJECT_ALIASES = {
    "ngoại ngữ tiếng anh": ["ngoại ngữ", "tiếng anh", "ngoại ngữ 1", "english"],
    "tn-xh": ["tự nhiên và xã hội", "tn-xh"],
    "tự nhiên và xã hội": ["tự nhiên và xã hội", "tn-xh"],
}


def set_merged_safe(ws, coord: str, value: str):
    for merged_range in ws.merged_cells.ranges:
        if coord in merged_range:
            ws.cell(merged_range.min_row, merged_range.min_col).value = value
            return
    ws[coord] = value


def set_cell_if_exists(ws, coord: str, value: str):
    set_merged_safe(ws, coord, value or "")


def replace_prefixed_cell(ws, coord: str, prefix: str, value: str):
    set_merged_safe(ws, coord, f"{prefix}{value or ''}")


def find_sheet(wb, keyword: str):
    k = norm_key(keyword)
    for ws in wb.worksheets:
        if k in norm_key(ws.title):
            return ws
    return None


def match_label(cell_value: str, desired: str) -> bool:
    cv = norm_key(cell_value).replace("\n", " ")
    de = norm_key(desired)
    if not cv or not de:
        return False
    if de in cv or cv in de:
        return True
    for alias_key, aliases in SUBJECT_ALIASES.items():
        if de == norm_key(alias_key) or de in [norm_key(a) for a in aliases]:
            return any(norm_key(a) in cv for a in aliases)
    return False


def fill_rows_by_first_col(
    ws,
    data_map: Dict[str, SubjectResult],
    level_col: str,
    score_col: Optional[str],
    comment_col: str,
    expected_keys: Optional[List[str]] = None,
):
    keys = expected_keys or list(data_map.keys())
    for r in range(1, ws.max_row + 1):
        label = norm_text(ws[f"A{r}"].value)
        if not label:
            continue
        for key in keys:
            if match_label(label, key):
                set_merged_safe(ws, f"{level_col}{r}", "")
                if score_col:
                    set_merged_safe(ws, f"{score_col}{r}", "")
                set_merged_safe(ws, f"{comment_col}{r}", "")
                val = data_map.get(key)
                if val is None:
                    val = next((item_val for item_key, item_val in data_map.items() if match_label(label, item_key)), SubjectResult())
                if val.level:
                    set_merged_safe(ws, f"{level_col}{r}", val.level)
                if score_col and val.score:
                    set_merged_safe(ws, f"{score_col}{r}", val.score)
                if val.comment:
                    set_merged_safe(ws, f"{comment_col}{r}", val.comment)
                break


def load_field_map(template_path: Path) -> Optional[dict]:
    csv_path = template_path.parent / "field_map.csv"
    if csv_path.exists():
        return load_field_map_csv(csv_path)

    json_path = template_path.parent / "field_map.json"
    if json_path.exists():
        return json.loads(json_path.read_text(encoding="utf-8"))
    return None


def load_field_map_csv(map_path: Path) -> dict:
    field_map = {"sheets": [], "tables": []}
    sheet_index = {}
    table_index = {}

    with map_path.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if not row or row.get("enabled", "1").strip() == "0":
                continue
            kind = row.get("type", "").strip().lower()
            name = row.get("name", "").strip()
            match_type = row.get("sheet_match_type", "").strip()
            match_value = row.get("sheet_match_value", "").strip()
            if not kind or not name or not match_type or not match_value:
                continue

            match = {match_type: match_value}
            if kind == "cell":
                sheet = sheet_index.get(name)
                if sheet is None:
                    sheet = {
                        "name": name,
                        "match": match,
                        "cells": [],
                    }
                    rename_template = row.get("rename_template", "").strip()
                    if rename_template:
                        sheet["rename"] = {"template": rename_template}
                    sheet_index[name] = sheet
                    field_map["sheets"].append(sheet)

                cell_map = {"cell": row.get("cell", "").strip()}
                for key in ("field", "prefix", "template", "transform"):
                    value = row.get(key, "") if key in ("prefix", "template") else row.get(key, "").strip()
                    if value:
                        cell_map[key] = value
                if cell_map["cell"]:
                    sheet["cells"].append(cell_map)

            elif kind == "table":
                table = table_index.get(name)
                if table is None:
                    columns = {}
                    for attr, csv_key in (("level", "level_col"), ("score", "score_col"), ("comment", "comment_col")):
                        value = row.get(csv_key, "").strip()
                        if value:
                            columns[attr] = value
                    table = {
                        "name": name,
                        "match": match,
                        "source": row.get("source", "").strip(),
                        "columns": columns,
                        "rows": [],
                    }
                    table_index[name] = table
                    field_map["tables"].append(table)

                table_row = {
                    "row": int(row.get("row", "0") or 0),
                    "key": row.get("key", "").strip(),
                }
                aliases = [x.strip() for x in row.get("aliases", "").split("|") if x.strip()]
                if aliases:
                    table_row["aliases"] = aliases
                if table_row["row"] and table_row["key"]:
                    table["rows"].append(table_row)

    return field_map


def match_sheet_by_rule(wb, rule: dict):
    exact = rule.get("exact")
    contains = rule.get("contains")
    starts_with = rule.get("starts_with")
    for ws in wb.worksheets:
        title = norm_key(ws.title)
        if exact and title == norm_key(exact):
            return ws
        if contains and norm_key(contains) in title:
            return ws
        if starts_with and title.startswith(norm_key(starts_with)):
            return ws
    return None


def get_data_field(data: HocBaData, field_name: str) -> str:
    if field_name == "qualities_comments":
        return "\n".join(
            data.qualities[key].comment
            for key in QUALITY_ITEMS
            if key in data.qualities and data.qualities[key].comment
        )
    if field_name == "common_competencies_comments":
        return "\n".join(
            data.common_competencies[key].comment
            for key in COMMON_COMP_ITEMS
            if key in data.common_competencies and data.common_competencies[key].comment
        )
    if field_name == "specific_competencies_comments":
        return "\n".join(
            data.specific_competencies[key].comment
            for key in SPECIFIC_COMP_ITEMS
            if key in data.specific_competencies and data.specific_competencies[key].comment
        )
    return str(getattr(data, field_name, "") or "")


def transform_value(value: str, transform: str) -> str:
    if transform == "school_year_spaced":
        return value.replace("-", " - ")
    return value


def render_cell_value(data: HocBaData, cell_map: dict) -> str:
    if "template" in cell_map:
        values = {field.name: getattr(data, field.name, "") or "" for field in data.__dataclass_fields__.values()}
        return cell_map["template"].format(**values)
    value = get_data_field(data, cell_map.get("field", ""))
    value = transform_value(value, cell_map.get("transform", ""))
    return f"{cell_map.get('prefix', '')}{value}"


def find_result_by_keys(data_map: Dict[str, SubjectResult], keys: List[str]) -> SubjectResult:
    for key in keys:
        if key in data_map:
            return data_map[key]
    for data_key, data_value in data_map.items():
        if any(match_label(data_key, key) or match_label(key, data_key) for key in keys):
            return data_value
    return SubjectResult()


def apply_field_map(wb, data: HocBaData, field_map: dict):
    for sheet_map in field_map.get("sheets", []):
        ws = match_sheet_by_rule(wb, sheet_map.get("match", {}))
        if not ws:
            continue
        for cell_map in sheet_map.get("cells", []):
            set_merged_safe(ws, cell_map["cell"], render_cell_value(data, cell_map))

        rename = sheet_map.get("rename")
        if rename:
            ws.title = rename["template"].format(
                class_name=data.class_name or "",
                school_year=data.school_year or "",
            )[:31]

    for table_map in field_map.get("tables", []):
        ws = match_sheet_by_rule(wb, table_map.get("match", {}))
        if not ws:
            continue
        data_map = getattr(data, table_map.get("source", ""), {}) or {}
        columns = table_map.get("columns", {})
        for row_map in table_map.get("rows", []):
            row = row_map["row"]
            keys = [row_map["key"], *row_map.get("aliases", [])]
            result = find_result_by_keys(data_map, keys)
            for attr, col in columns.items():
                set_merged_safe(ws, f"{col}{row}", getattr(result, attr, "") or "")


def is_thcs_template(wb) -> bool:
    titles = {norm_key(ws.title) for ws in wb.worksheets}
    return "table 2" in titles and "table 3" in titles


def sheet_by_exact_title(wb, title: str):
    title_key = norm_key(title)
    for ws in wb.worksheets:
        if norm_key(ws.title) == title_key:
            return ws
    return None


def thcs_text_with_prefix(prefix: str, value: str) -> str:
    return f"{prefix}{value or ''}"


def build_placeholder_index(wb) -> Dict[int, List[Tuple[object, str]]]:
    index: Dict[int, List[Tuple[object, str]]] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                for match in re.finditer(r"\[(\d+)\]", cell.value):
                    index.setdefault(int(match.group(1)), []).append((ws, cell.coordinate))
    return index


def set_placeholder(index: Dict[int, List[Tuple[object, str]]], number: int, value: str, sheet_name: str = ""):
    targets = index.get(number, [])
    if sheet_name:
        sheet_key = norm_key(sheet_name)
        targets = [target for target in targets if norm_key(target[0].title) == sheet_key]
    if not targets:
        return
    ws, coord = targets[0]
    set_merged_safe(ws, coord, value or "")


def fill_thcs_excel(wb, data: HocBaData):
    placeholders = build_placeholder_index(wb)

    # Trang bìa
    set_placeholder(placeholders, 1, data.school)
    set_placeholder(placeholders, 2, thcs_text_with_prefix("Xã/Phường/Đặc khu: ", data.ward))
    set_placeholder(placeholders, 3, thcs_text_with_prefix("Tỉnh/Thành phố: ", data.province))
    set_placeholder(placeholders, 4, data.student_name.upper() if data.student_name else "")
    register = data.school_register_no or "………………………………"
    register_text = register if register.endswith("/THCS") else f"{register}/THCS"
    set_placeholder(placeholders, 5, f"Số sổ đăng bộ PCGD: {register_text}")

    # Trang thông tin học sinh và quá trình học tập.
    for number, value in (
        (6, data.student_name),
        (7, data.gender),
        (8, data.birth_date),
        (9, data.ethnicity),
        (10, data.birth_place),
        (11, data.policy_object),
        (12, data.address),
        (13, data.father),
        (14, data.father_job),
        (15, data.mother),
        (16, data.mother_job),
        (17, data.guardian),
        (18, data.guardian_job),
        (19, data.issue_date),
        (20, data.principal_title or "HIỆU TRƯỞNG"),
        (21, data.principal_subtitle),
        (22, data.principal),
    ):
        set_placeholder(placeholders, number, value)

    for idx in range(4):
        record = data.study_records[idx] if idx < len(data.study_records) else StudyRecord()
        base = 23 + idx * 3
        set_placeholder(placeholders, base, record.school_year)
        set_placeholder(placeholders, base + 1, record.class_name)
        set_placeholder(placeholders, base + 2, record.school)

    # Bảng điểm môn học.
    set_placeholder(placeholders, 35, f"Họ và tên:   {data.student_name}")
    set_placeholder(placeholders, 36, f"Lớp: {data.class_name}")
    set_placeholder(placeholders, 37, f"Năm học: {data.school_year}")
    subject_keys = [
        "Toán học",
        "Lịch sử và Địa lí",
        "Khoa học tự nhiên",
        "Tin học",
        "Ngữ văn",
        "Ngoại ngữ 1 (Tiếng Anh)",
        "GDCD",
        "Công nghệ",
        "Giáo dục thể chất",
        "Nghệ thuật",
        "Ngoại ngữ 2",
        "Nội dung giáo dục của địa phương",
        "Hoạt động trải nghiệm, hướng nghiệp",
    ]
    subject_placeholders = {
        "Toán học": 38,
        "Lịch sử và Địa lí": 44,
        "Khoa học tự nhiên": 50,
        "Tin học": 56,
        "Ngữ văn": 62,
        "Ngoại ngữ 1 (Tiếng Anh)": 68,
        "GDCD": 74,
        "Công nghệ": 80,
        "Giáo dục thể chất": 86,
        "Nghệ thuật": 92,
        "Ngoại ngữ 2": 132,
        "Nội dung giáo dục của địa phương": 98,
        "Hoạt động trải nghiệm, hướng nghiệp": 104,
    }
    for key in subject_keys:
        result = data.thcs_subjects.get(key, ThcsSubjectResult())
        base = subject_placeholders[key]
        set_placeholder(placeholders, base, key, "Table 2")
        set_placeholder(placeholders, base + 1, result.hk1, "Table 2")
        set_placeholder(placeholders, base + 2, result.hk2, "Table 2")
        set_placeholder(placeholders, base + 3, result.year, "Table 2")
        set_placeholder(placeholders, base + 4, result.retake, "Table 2")
        set_placeholder(placeholders, base + 5, result.teacher, "Table 2")
    set_placeholder(placeholders, 110, "Trong trang này có sửa chữa ở không chỗ, thuộc môn học, hoạt động giáo dục: ...................................")
    set_placeholder(placeholders, 111, data.teacher)
    set_placeholder(placeholders, 112, data.principal)

    # Bảng tổng hợp rèn luyện, học tập và nhận xét cuối năm.
    set_placeholder(placeholders, 113, f"Họ tên học sinh: {data.student_name}                         Lớp: {data.class_name}          Năm học: {data.school_year}", "Table 3")
    set_placeholder(placeholders, 116, data.promotion, "Table 3")
    for idx, key in enumerate(("hk1", "hk2", "year")):
        result = data.summary_results.get(key, SummaryResult())
        base = 117 + idx * 3
        set_placeholder(placeholders, base, result.conduct, "Table 3")
        set_placeholder(placeholders, base + 1, result.learning, "Table 3")
        set_placeholder(placeholders, base + 2, result.absent, "Table 3")

    completion = data.completion or "......................................................................................................................................................."
    certificate = data.certificate or "................................................................. Loại: ............................................"
    contest = data.contest_result or "..........................................................................................................................................................."
    reward = data.reward or "..................................................................................................................."
    set_placeholder(
        placeholders,
        126,
        "Nếu là lớp cuối cấp thì ghi Hoàn thành hay không hoàn thành chương trình trung học cơ sở:\n"
        f"{completion}\n"
        f"- Chứng chỉ (nếu có): {certificate}\n"
        f"- Kết quả tham gia các cuộc thi (nếu có): {contest}\n"
        f"- Khen thưởng (nếu có): {reward}",
    )
    set_placeholder(placeholders, 127, data.homeroom_comment)
    set_placeholder(placeholders, 129, data.teacher)
    set_placeholder(placeholders, 130, data.sign_date)
    set_placeholder(placeholders, 131, data.principal)
    set_placeholder(placeholders, 132, data.sign_page_principal_title or "Hiệu trưởng")


def fill_hoc_ba_excel(template_path: Path, output_path: Path, data: HocBaData):
    wb = load_workbook(template_path)
    if is_thcs_template(wb):
        fill_thcs_excel(wb, data)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return

    field_map = load_field_map(template_path)
    if field_map:
        apply_field_map(wb, data, field_map)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return

    # Bìa học bạ
    ws = find_sheet(wb, "Bìa")
    if ws:
        replace_prefixed_cell(ws, "A36", "Họ và tên học sinh: ", data.student_name)
        replace_prefixed_cell(ws, "A38", "Trường: ", data.school)
        replace_prefixed_cell(ws, "A40", "Xã (Phường, Thị trấn): ", data.ward)
        replace_prefixed_cell(ws, "A42", "Huyện (Thành phố, Quận, Thị xã): ", data.district)
        replace_prefixed_cell(ws, "A44", "Tỉnh (Thành phố): ", data.province)

    # Thông tin học sinh
    ws = find_sheet(wb, "Thông tin")
    if ws:
        replace_prefixed_cell(ws, "A4", "Họ và tên học sinh: ", data.student_name)
        replace_prefixed_cell(ws, "P4", "Giới tính: ", data.gender)
        replace_prefixed_cell(ws, "A5", "Ngày, tháng, năm sinh: ", data.birth_date)
        replace_prefixed_cell(ws, "K5", "Dân tộc: ", data.ethnicity)
        replace_prefixed_cell(ws, "P5", "Quốc tịch: ", data.nationality)
        replace_prefixed_cell(ws, "A6", "Nơi sinh: ", data.birth_place)
        replace_prefixed_cell(ws, "A7", "Quê quán: ", data.hometown)
        replace_prefixed_cell(ws, "A8", "Nơi ở hiện nay: ", data.address)
        replace_prefixed_cell(ws, "A9", "Họ và tên cha: ", data.father)
        replace_prefixed_cell(ws, "A10", "Họ và tên mẹ: ", data.mother)
        replace_prefixed_cell(ws, "A11", "Người giám hộ (nếu có): ", data.guardian)
        ws["K14"] = data.issue_date or ""
        set_cell_if_exists(ws, "K21", data.principal)
        ws["A28"] = data.school_year.replace("-", " - ") if data.school_year else ""
        set_cell_if_exists(ws, "D28", data.class_name)
        set_cell_if_exists(ws, "F28", data.school)
        set_cell_if_exists(ws, "Q28", data.entry_date)

    # Môn học + Phẩm chất
    ws = None
    for s in wb.worksheets:
        if norm_key(s.title).startswith("mh"):
            ws = s
            break
    if ws:
        replace_prefixed_cell(ws, "A1", "Họ và tên học sinh: ", data.student_name)
        replace_prefixed_cell(ws, "O1", "Lớp: ", data.class_name)
        replace_prefixed_cell(ws, "A2", "Chiều cao: ", data.height)
        replace_prefixed_cell(ws, "O2", "Cân nặng: ", data.weight)
        replace_prefixed_cell(ws, "A3", "Số ngày nghỉ có phép: ", data.absent_allowed)
        replace_prefixed_cell(ws, "O3", "Số ngày nghỉ không phép: ", data.absent_not_allowed)
        fill_rows_by_first_col(
            ws,
            data.subjects,
            level_col="F",
            score_col="H",
            comment_col="J",
            expected_keys=[key for key, _ in SUBJECT_ITEMS],
        )
        fill_rows_by_first_col(
            ws,
            data.qualities,
            level_col="G",
            score_col=None,
            comment_col="J",
            expected_keys=QUALITY_ITEMS,
        )
        # đổi tên sheet gọn hơn nếu có lớp/năm học
        new_title = f"MH {data.class_name} ({data.school_year})"[:31] if data.class_name else ws.title
        ws.title = new_title

    # Năng lực + kết quả cuối năm
    ws = None
    for s in wb.worksheets:
        if norm_key(s.title).startswith("nl"):
            ws = s
            break
    if ws:
        ws["A1"] = data.school or ""
        ws["O1"] = f"Năm học {data.school_year}" if data.school_year else ""
        fill_rows_by_first_col(
            ws,
            data.common_competencies,
            level_col="G",
            score_col=None,
            comment_col="J",
            expected_keys=COMMON_COMP_ITEMS,
        )
        fill_rows_by_first_col(
            ws,
            data.specific_competencies,
            level_col="G",
            score_col=None,
            comment_col="J",
            expected_keys=SPECIFIC_COMP_ITEMS,
        )
        ws["A20"] = f"4. Đánh giá kết quả giáo dục:  {data.result or ''}"
        ws["A22"] = f"5. Khen thưởng: {data.reward or ''}"
        ws["A29"] = data.completion or ""
        ws["K32"] = data.sign_date or ""
        set_cell_if_exists(ws, "A40", data.principal)
        set_cell_if_exists(ws, "K40", data.teacher)
        new_title = f"NL, PC {data.class_name} ({data.school_year})"[:31] if data.class_name else ws.title
        ws.title = new_title

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def convert_one_pdf(pdf_path: Path, template_path: Path, output_dir: Path) -> Path:
    data = parse_hoc_ba_pdf(pdf_path)
    class_part = safe_filename(data.class_name or "CHUA_CO_LOP")
    name_part = safe_filename(data.student_name or pdf_path.stem)
    output_path = output_dir / f"{class_part}_{name_part}.xlsx"
    fill_hoc_ba_excel(template_path, output_path, data)
    return output_path


# =========================
# 5) GIAO DIỆN TKINTER
# =========================

class HocBaApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Chuyển PDF học bạ sang Excel mẫu - Võ Thành Khải")
        self.geometry("1050x650")
        self.minsize(950, 560)

        self.input_dir = tk.StringVar(value=str(app_base_dir() / "input_sample"))
        self.output_dir = tk.StringVar(value=str(app_base_dir() / "output"))
        default_template = app_base_dir() / "templates" / "hoc_ba_mau (2).xlsx"
        if not default_template.exists():
            default_template = app_base_dir() / "templates" / "hoc_ba_mau.xlsx"
        self.template_file = tk.StringVar(value=str(default_template) if default_template.exists() else "")
        self.status_text = tk.StringVar(value="Sẵn sàng")
        self.items: Dict[str, Dict[str, object]] = {}

        self._build_ui()
        if Path(self.input_dir.get()).exists():
            self.load_pdf_files()

    def _build_ui(self):
        frm = ttk.Frame(self, padding=10)
        frm.pack(fill="both", expand=True)

        # đường dẫn
        path_frame = ttk.LabelFrame(frm, text="Thư mục và file mẫu", padding=8)
        path_frame.pack(fill="x")

        self._path_row(path_frame, 0, "Thư mục input PDF:", self.input_dir, self.choose_input_dir)
        self._path_row(path_frame, 1, "Thư mục output Excel:", self.output_dir, self.choose_output_dir)
        self._path_row(path_frame, 2, "File Excel mẫu:", self.template_file, self.choose_template_file)

        # nút chức năng
        btn_frame = ttk.Frame(frm)
        btn_frame.pack(fill="x", pady=(10, 6))
        ttk.Button(btn_frame, text="Nạp danh sách PDF", command=self.load_pdf_files).pack(side="left", padx=3)
        ttk.Button(btn_frame, text="Check chọn tất cả", command=self.check_all).pack(side="left", padx=3)
        ttk.Button(btn_frame, text="Bỏ check tất cả", command=self.uncheck_all).pack(side="left", padx=3)
        ttk.Button(btn_frame, text="Xóa list file đã chọn", command=self.remove_checked).pack(side="left", padx=3)
        ttk.Button(btn_frame, text="Thực hiện chuyển file đã chọn", command=self.start_convert).pack(side="right", padx=3)
        if _HAS_UPDATER:
            ttk.Button(btn_frame, text="Kiểm tra cập nhật", command=self.check_for_update).pack(side="right", padx=3)

        # grid
        grid_frame = ttk.LabelFrame(frm, text="Danh sách file input", padding=6)
        grid_frame.pack(fill="both", expand=True)
        cols = ("checked", "file", "student", "class", "status")
        self.tree = ttk.Treeview(grid_frame, columns=cols, show="headings", selectmode="extended")
        self.tree.heading("checked", text="Chọn")
        self.tree.heading("file", text="File PDF")
        self.tree.heading("student", text="Tên học sinh")
        self.tree.heading("class", text="Lớp")
        self.tree.heading("status", text="Trạng thái")
        self.tree.column("checked", width=70, anchor="center")
        self.tree.column("file", width=360)
        self.tree.column("student", width=230)
        self.tree.column("class", width=90, anchor="center")
        self.tree.column("status", width=230)
        self.tree.pack(side="left", fill="both", expand=True)
        self.tree.bind("<Double-1>", self.toggle_selected)
        self.tree.bind("<space>", self.toggle_selected)

        yscroll = ttk.Scrollbar(grid_frame, orient="vertical", command=self.tree.yview)
        yscroll.pack(side="right", fill="y")
        self.tree.configure(yscrollcommand=yscroll.set)

        # log/status
        bottom = ttk.Frame(frm)
        bottom.pack(fill="x", pady=(8, 0))
        ttk.Label(bottom, textvariable=self.status_text).pack(side="left")
        ttk.Label(bottom, text="Tác giả: Võ Thành Khải - 0913046881 / 0913784333").pack(side="right")

    def _path_row(self, parent, row, label, var, cmd):
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 6), pady=3)
        ttk.Entry(parent, textvariable=var).grid(row=row, column=1, sticky="ew", pady=3)
        ttk.Button(parent, text="Chọn...", command=cmd).grid(row=row, column=2, padx=(6, 0), pady=3)
        parent.columnconfigure(1, weight=1)

    def choose_input_dir(self):
        d = filedialog.askdirectory(title="Chọn thư mục chứa file PDF")
        if d:
            self.input_dir.set(d)
            self.load_pdf_files()

    def choose_output_dir(self):
        d = filedialog.askdirectory(title="Chọn thư mục lưu file Excel kết quả")
        if d:
            self.output_dir.set(d)

    def choose_template_file(self):
        f = filedialog.askopenfilename(title="Chọn file Excel mẫu", filetypes=[("Excel", "*.xlsx")])
        if f:
            self.template_file.set(f)

    def load_pdf_files(self):
        folder = Path(self.input_dir.get())
        if not folder.exists():
            messagebox.showwarning("Chưa có thư mục", "Thư mục input không tồn tại")
            return
        self.tree.delete(*self.tree.get_children())
        self.items.clear()
        files = sorted(folder.glob("*.pdf"))
        for p in files:
            iid = str(p)
            student = ""
            cls = ""
            try:
                data = parse_hoc_ba_pdf(p)
                student = data.student_name
                cls = data.class_name
            except Exception:
                student = "Chưa đọc nhanh được"
            self.items[iid] = {"path": p, "checked": True}
            self.tree.insert("", "end", iid=iid, values=("☑", p.name, student, cls, "Chờ chuyển"))
        self.status_text.set(f"Đã nạp {len(files)} file PDF")

    def toggle_selected(self, event=None):
        selected = self.tree.selection()
        if not selected:
            item = self.tree.identify_row(event.y) if event else None
            selected = [item] if item else []
        for iid in selected:
            if iid in self.items:
                self.items[iid]["checked"] = not bool(self.items[iid]["checked"])
                self._refresh_item_check(iid)

    def _refresh_item_check(self, iid: str):
        vals = list(self.tree.item(iid, "values"))
        vals[0] = "☑" if self.items[iid]["checked"] else "☐"
        self.tree.item(iid, values=vals)

    def check_all(self):
        for iid in self.items:
            self.items[iid]["checked"] = True
            self._refresh_item_check(iid)

    def uncheck_all(self):
        for iid in self.items:
            self.items[iid]["checked"] = False
            self._refresh_item_check(iid)

    def remove_checked(self):
        remove = [iid for iid, info in self.items.items() if info.get("checked")]
        for iid in remove:
            self.items.pop(iid, None)
            try:
                self.tree.delete(iid)
            except Exception:
                pass
        self.status_text.set(f"Đã xóa khỏi list {len(remove)} file được chọn")

    def set_status_for_item(self, iid: str, status: str):
        vals = list(self.tree.item(iid, "values"))
        if vals:
            vals[4] = status
            self.tree.item(iid, values=vals)

    def check_for_update(self):
        self.status_text.set("Đang kiểm tra cập nhật...")

        def run():
            try:
                updater = AutoUpdater()
                updater.check_and_update(show_dialog=True)
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Lỗi", f"Không thể kiểm tra cập nhật:\n{e}"))
            finally:
                self.after(0, lambda: self.status_text.set("Sẵn sàng"))

        threading.Thread(target=run, daemon=True).start()

    def start_convert(self):
        selected = [(iid, info) for iid, info in self.items.items() if info.get("checked")]
        if not selected:
            messagebox.showinfo("Chưa chọn file", "Vui lòng check chọn ít nhất 1 file PDF")
            return
        template = Path(self.template_file.get())
        output = Path(self.output_dir.get())
        if not template.exists():
            messagebox.showerror("Thiếu file mẫu", "File Excel mẫu không tồn tại")
            return
        output.mkdir(parents=True, exist_ok=True)
        threading.Thread(target=self._convert_thread, args=(selected, template, output), daemon=True).start()

    def _convert_thread(self, selected, template: Path, output: Path):
        ok = 0
        fail = 0
        for idx, (iid, info) in enumerate(selected, start=1):
            p: Path = info["path"]
            self.after(0, lambda iid=iid: self.set_status_for_item(iid, "Đang chuyển..."))
            try:
                out = convert_one_pdf(p, template, output)
                ok += 1
                self.after(0, lambda iid=iid, out=out: self.set_status_for_item(iid, f"OK: {out.name}"))
            except Exception as e:
                fail += 1
                err_file = output / f"LOI_{safe_filename(p.stem)}.txt"
                err_file.write_text(traceback.format_exc(), encoding="utf-8")
                self.after(0, lambda iid=iid, e=e: self.set_status_for_item(iid, f"Lỗi: {e}"))
            self.after(0, lambda idx=idx, total=len(selected), ok=ok, fail=fail: self.status_text.set(f"Đang xử lý {idx}/{total} | OK {ok} | Lỗi {fail}"))
        self.after(0, lambda: messagebox.showinfo("Hoàn tất", f"Đã chuyển xong. OK: {ok}, Lỗi: {fail}\nThư mục output: {output}"))
        self.after(0, lambda: self.status_text.set(f"Hoàn tất | OK {ok} | Lỗi {fail}"))


if __name__ == "__main__":
    if ttk is None:
        raise SystemExit("Tkinter is not available. Run the web version with: python web_app.py")
    app = HocBaApp()
    app.mainloop()
