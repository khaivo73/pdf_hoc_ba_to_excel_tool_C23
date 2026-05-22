import openpyxl
import re

# Load template XLSX
template_file = r'templates\hoc_ba_mau.xlsx'
wb = openpyxl.load_workbook(template_file)

print("=== TRÍCH XUẤT CÁC TRƯỜNG THÔNG TIN ===\n")

fields_found = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    fields_found[sheet_name] = []
    
    print(f"\n--- Sheet: {sheet_name} ---")
    
    # Iterate through all cells to find numbered fields like [1], [2], etc.
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                # Look for pattern [number]
                matches = re.findall(r'\[(\d+)\]', cell.value)
                if matches:
                    field_num = matches[0]
                    fields_found[sheet_name].append({
                        'number': field_num,
                        'text': cell.value,
                        'cell': cell.coordinate
                    })
                    print(f"[{field_num}] {cell.value} (Cell: {cell.coordinate})")

print("\n\n=== TÓM TẮT CÁC TRƯỜNG ĐƯỢC ĐÁNH SỐ ===")
for sheet, fields in fields_found.items():
    if fields:
        print(f"\n{sheet}:")
        for f in sorted(fields, key=lambda x: int(x['number'])):
            print(f"  [{f['number']}] {f['text']}")
